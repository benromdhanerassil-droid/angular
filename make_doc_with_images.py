import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def generate_final_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("COMPTE RENDU DES ATELIERS ANGULAR (1, 2 & 3)")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(26, 82, 118)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Application : Gestion des Watches (MesWatches)\nFramework : Angular 17+ (Standalone Components, Control Flow, Services & Routing)")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(86, 101, 115)

    doc.add_paragraph()

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(31, 78, 121)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(46, 117, 182)

    def add_code(text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F4F6F6")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text.strip())
        r.font.name = 'Consolas'
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(33, 47, 60)
        doc.add_paragraph()

    def add_image_block(image_path, caption):
        if os.path.exists(image_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            doc.add_picture(image_path, width=Inches(5.8))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(f"Figure : {caption}")
            r_cap.font.name = 'Arial'
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(86, 101, 115)

    shots_dir = r"c:\Users\MSI\MesProduits\screenshots"
    shot_liste = os.path.join(shots_dir, "01_liste_watches.png")
    shot_ajout = os.path.join(shots_dir, "02_ajout_watche.png")
    shot_modif = os.path.join(shots_dir, "03_modif_watche.png")

    # --- Atelier 1 ---
    add_h1("1. ATELIER 01 : Initialisation, Composants & Navbar Bootstrap")
    add_h2("1.1. Commandes Angular CLI")
    add_code("""ng new MesWatches
npm install bootstrap@5.3.3
ng g c watches --skip-tests
ng g c add-watche --skip-tests""")

    add_h2("1.2. Configuration de Bootstrap (angular.json)")
    add_code(""""styles": [
  "src/styles.css",
  "node_modules/bootstrap/dist/css/bootstrap.min.css"
],
"scripts": [
  "node_modules/bootstrap/dist/js/bootstrap.bundle.min.js"
]""")

    add_h2("1.3. Barre de navigation (app.html)")
    add_code("""<nav class="navbar navbar-expand-lg navbar-light bg-light">
  <div class="container-fluid">
    <a class="navbar-brand" href="#">Gestion des Watches</a>
    <div class="collapse navbar-collapse" id="navbarSupportedContent">
      <ul class="navbar-nav me-auto mb-2 mb-lg-0">
        <li class="nav-item">
          <a class="nav-link active" routerLink="/watches">Home</a>
        </li>
        <li class="nav-item dropdown">
          <a class="nav-link dropdown-toggle" href="#" id="navbarDropdownWatches"
            role="button" data-bs-toggle="dropdown">Watches</a>
          <ul class="dropdown-menu">
            <li><a class="dropdown-item" routerLink="/add-watche">Ajouter</a></li>
            <li><a class="dropdown-item" routerLink="/watches">Lister</a></li>
          </ul>
        </li>
      </ul>
    </div>
  </div>
</nav>
<router-outlet />""")

    # --- Atelier 2 ---
    add_h1("2. ATELIER 02 : Modèle Watche, Service, $index et Formulaire d'Ajout")
    add_h2("2.1. Modèle de données (src/app/model/watche.model.ts)")
    add_code("""export class Watche {
  idWatche?: number;
  marque?: string;
  prixWatche?: number;
  dateFabrication?: Date;
  typeMouvement?: string;
}""")

    add_h2("2.2. Service Watche (src/app/services/watche.service.ts)")
    add_code("""import { Injectable } from '@angular/core';
import { Watche } from '../model/watche.model';

@Injectable({
  providedIn: 'root'
})
export class WatcheService {
  watches: Watche[];
  watche!: Watche;

  constructor() {
    this.watches = [
      { idWatche: 1, marque: 'Rolex', prixWatche: 12500, dateFabrication: new Date('2021-03-15'), typeMouvement: 'automatique' },
      { idWatche: 2, marque: 'Seiko', prixWatche: 350.5, dateFabrication: new Date('2019-07-22'), typeMouvement: 'quartz' },
      { idWatche: 3, marque: 'Omega', prixWatche: 5800, dateFabrication: new Date('2020-11-08'), typeMouvement: 'mécanique' }
    ];
  }

  listeWatches(): Watche[] { return this.watches; }
  ajouterWatche(w: Watche) { this.watches.push(w); }
}""")

    add_h2("2.3. Affichage du tableau avec @for et $index (watches.component.html)")
    add_code("""<table class="table table-striped">
  <tr>
    <th>N°</th>
    <th>Id</th>
    <th>Marque</th>
    <th>Prix</th>
    <th>Date Fabrication</th>
    <th>Type Mouvement</th>
    <th>Actions</th>
  </tr>
  @for (watche of watches; track watche.idWatche; let index = $index) {
    <tbody>
      <tr>
        <td>{{index}}</td>
        <td>{{watche.idWatche}}</td>
        <td>{{watche.marque}}</td>
        <td>{{watche.prixWatche}}</td>
        <td>{{watche.dateFabrication | date:'dd/MM/yyyy'}}</td>
        <td>{{watche.typeMouvement}}</td>
        <td>
          <a class="btn btn-danger" (click)="supprimerWatche(watche)">Supprimer</a>
          <a class="btn btn-success" [routerLink]="['/updateWatche', watche.idWatche]">Modifier</a>
        </td>
      </tr>
    </tbody>
  }
</table>""")
    add_image_block(shot_liste, "Affichage de la liste des watches avec la colonne N° ($index), tableau Bootstrap et actions")

    add_h2("2.4. Formulaire d'ajout de Watche (add-watche.component.html)")
    add_image_block(shot_ajout, "Formulaire d'ajout d'une nouvelle watche avec ngModel")

    # --- Atelier 3 ---
    add_h1("3. ATELIER 03 : Modification, Suppression & ActivatedRoute")
    add_h2("3.1. Méthodes CRUD dans WatcheService")
    add_code("""supprimerWatche(w: Watche) {
  const index = this.watches.indexOf(w, 0);
  if (index > -1) { this.watches.splice(index, 1); }
}

consulterWatche(id: number): Watche {
  this.watche = this.watches.find(w => w.idWatche == id)!;
  return this.watche;
}

updateWatche(w: Watche) {
  let index = this.watches.indexOf(w, 0);
  if (index === -1) { index = this.watches.findIndex(item => item.idWatche == w.idWatche); }
  if (index > -1) {
    this.watches.splice(index, 1);
    this.watches.splice(index, 0, w);
  }
}""")

    add_h2("3.2. Formulaire de modification (update-watche.component.html)")
    add_image_block(shot_modif, "Formulaire de modification d'une watche avec ID en lecture seule (readonly)")

    add_h2("3.3. Configuration des Routes finales (app.routes.ts)")
    add_code("""import { Routes } from '@angular/router';
import { WatchesComponent } from './watches/watches';
import { AddWatcheComponent } from './add-watche/add-watche';
import { UpdateWatcheComponent } from './update-watche/update-watche';

export const routes: Routes = [
  { path: 'watches', component: WatchesComponent },
  { path: 'add-watche', component: AddWatcheComponent },
  { path: '', redirectTo: 'watches', pathMatch: 'full' },
  { path: 'updateWatche/:id', component: UpdateWatcheComponent }
];""")

    output_path = r"c:\Users\MSI\MesProduits\Compte_Rendu_Ateliers_Watches.docx"
    doc.save(output_path)
    print("FINISHED: Docx with embedded images saved to", output_path)

if __name__ == '__main__':
    generate_final_report()
